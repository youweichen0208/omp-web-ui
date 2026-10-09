"""Versioned, local-only Docling worker. stdout is one JSON response; logs use stderr."""
import hashlib
import importlib.metadata
import json
import logging
import os
from pathlib import Path
import platform
import re
import shutil
import sys
import tempfile
from contextlib import redirect_stdout
from concurrent.futures import ThreadPoolExecutor

VERSION = "2.136.0"
PROFILE = "docling-cpu-zh-v1"
MODEL_REVISIONS = {
	"docling-project/docling-layout-heron": "8f39ad3c0b4c58e9c2d2c84a38465abf757272d8",
	"docling-project/docling-models": "fc0f2d45e2218ea24bce5045f58a389aed16dc23",
	"docling-project/CodeFormulaV2": "ecedbe111d15c2dc60bfd4a823cbe80127b58af4",
}
logging.basicConfig(stream=sys.stderr, level=logging.WARNING)
# Hugging Face's HTTP downloader respects the host's proxy configuration and can
# resume downloads without requiring Xet's separate network transport.
os.environ.setdefault("HF_HUB_DISABLE_XET", "1")


def emit(value):
	print(json.dumps(value, ensure_ascii=False), flush=True)


def sha256(path):
	hash_value = hashlib.sha256()
	with path.open("rb") as stream:
		for chunk in iter(lambda: stream.read(1024 * 1024), b""):
			hash_value.update(chunk)
	return hash_value.hexdigest()


def parser_assets():
	return {path.name: sha256(path) for path in sorted(Path(__file__).parent.glob("*.py"))}


LOADED_ASSET_HASHES = parser_assets()


def offline():
	for name in ("HF_HUB_OFFLINE", "TRANSFORMERS_OFFLINE", "HF_DATASETS_OFFLINE", "HF_HUB_DISABLE_TELEMETRY"):
		os.environ[name] = "1"
	# Protect all optional backends as well as Hugging Face. Conversion never opens
	# TCP connections, including those to loopback or inherited proxy endpoints.
	def audit(event, args):
		if event in ("socket.connect", "socket.getaddrinfo", "socket.sendto"):
			raise PermissionError("Network access is disabled during document conversion")
	sys.addaudithook(audit)


def doctor(root):
	missing = []
	warnings = []
	try:
		actual = importlib.metadata.version("docling")
		if actual != VERSION:
			missing.append("Expected Docling " + VERSION + "; found " + actual)
	except importlib.metadata.PackageNotFoundError:
		actual = None
		missing.append("Docling is not installed")
	if sys.version_info[:2] != (3, 12):
		missing.append("This runtime profile requires Python 3.12")
	try:
		receipt = json.loads((root / "ready.json").read_text("utf-8"))
		if receipt.get("profile") != PROFILE or receipt.get("docling") != VERSION or not receipt.get("smokePassed"):
			missing.append("Runtime qualification receipt is invalid")
		if receipt.get("parserAssets") != parser_assets():
			missing.append("Parser implementation changed; run /pdf-md setup to requalify the offline profile")
		for relative, metadata in receipt.get("models", {}).items():
			path = root / "models" / relative
			if not path.is_file() or path.stat().st_size != metadata["size"] or sha256(path) != metadata["sha256"]:
				missing.append("Missing or changed model: " + relative)
		for name, version in receipt.get("packages", {}).items():
			try:
				if importlib.metadata.version(name) != version:
					missing.append("Runtime dependency changed: " + name)
			except importlib.metadata.PackageNotFoundError:
				missing.append("Runtime dependency missing: " + name)
		if not receipt.get("models"):
			missing.append("No model artifacts were recorded")
	except (OSError, ValueError, KeyError):
		missing.append("Run /pdf-md setup to download models and qualify offline conversion")
	return {"ready": not missing, "parserVersion": actual, "missing": missing, "warnings": warnings}


def converter(root):
	from docling.datamodel.accelerator_options import AcceleratorDevice, AcceleratorOptions
	from docling.datamodel.base_models import InputFormat
	from docling.datamodel.pipeline_options import (
		CodeFormulaVlmOptions, OcrMode, PdfPipelineOptions, RapidOcrOptions,
		TableFormerMode, TableStructureOptions,
	)
	from docling.datamodel.vlm_engine_options import TransformersVlmEngineOptions
	from docling.document_converter import DocumentConverter, PdfFormatOption
	options = PdfPipelineOptions(
		artifacts_path=root / "models",
		enable_remote_services=False,
		allow_external_plugins=False,
		accelerator_options=AcceleratorOptions(device=AcceleratorDevice.CPU, num_threads=4),
		do_ocr=True,
		ocr_options=RapidOcrOptions(backend="onnxruntime", lang=["ch"], model_size="small", mode=OcrMode.PDF_AWARE_LAYOUT_REGIONS),
		do_table_structure=True,
		table_structure_options=TableStructureOptions(mode=TableFormerMode.ACCURATE, do_cell_matching=True),
		do_code_enrichment=True,
		do_formula_enrichment=True,
		code_formula_options=CodeFormulaVlmOptions.from_preset("codeformulav2", engine_options=TransformersVlmEngineOptions(
			device=AcceleratorDevice.CPU, torch_dtype="float32", quantized=False, compile_model=False,
		)),
		generate_picture_images=True,
		generate_page_images=True,
	)
	options.heading_hierarchy_options.enabled = True
	return DocumentConverter(allowed_formats=[InputFormat.PDF, InputFormat.DOCX, InputFormat.PPTX, InputFormat.XLSX], format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=options)})


def extract_blocks(document, suffix):
	from docling_core.types.doc import TableItem
	blocks = []
	head = None
	for item, level in document.iterate_items():
		label = str(getattr(item, "label", ""))
		text = getattr(item, "text", "")
		if isinstance(item, TableItem):
			text = item.export_to_markdown(doc=document)
		if not text:
			continue
		if "section_header" in label or "title" in label:
			head = text
		locator = {"heading": head} if head else {}
		provenance = [p.model_dump(mode="json") for p in getattr(item, "prov", [])]
		if provenance:
			page = provenance[0].get("page_no")
			if page:
				locator["slide" if suffix == ".pptx" else "page"] = page
		blocks.append({"id": getattr(item, "self_ref", "block-" + str(len(blocks) + 1)), "text": text, "locator": locator, "provenance": provenance, "kind": label})
	return blocks


def excel_provenance(source, warnings):
	from openpyxl import load_workbook
	formulas = load_workbook(source, read_only=True, data_only=False)
	values = load_workbook(source, read_only=True, data_only=True)
	blocks = []
	metadata = []
	try:
		for sheet in formulas:
			cached = values[sheet.title]
			for row, cached_row in zip(sheet.iter_rows(), cached.iter_rows()):
				for cell, cached_cell in zip(row, cached_row):
					if cell.value is None:
						continue
					value = cached_cell.value
					is_formula = cell.data_type == "f"
					if is_formula:
						metadata.append({"sheet": sheet.title, "cell": cell.coordinate, "formula": str(cell.value), "cachedValue": value})
						if value is None:
							warnings.append("Missing formula cache: " + sheet.title + "!" + cell.coordinate)
					blocks.append({"id": "cell-" + str(len(blocks) + 1), "text": str(value) if value is not None else "[formula result unavailable]", "locator": {"sheet": sheet.title, "cell": cell.coordinate}})
		if metadata:
			warnings.append("Spreadsheet formula results are saved cached values; no formulas were evaluated and cache freshness cannot be verified.")
	finally:
		formulas.close()
		values.close()
	return blocks, metadata


def convert(root, request, engine=None):
	from docling_core.types.doc import ImageRefMode, TableItem
	source, output = Path(request["inputPath"]), Path(request["outputDir"])
	output.mkdir(parents=True, exist_ok=True)
	assets = output / "assets"
	assets.mkdir(exist_ok=True)
	warnings = []
	enrichment_failed = False
	material_defect = False
	native_verified_ids = set()
	expected_pages = None
	if source.suffix.lower() == ".pdf":
		import pypdfium2
		pdf = pypdfium2.PdfDocument(source)
		try:
			expected_pages = len(pdf)
			native = []
			for index in range(expected_pages):
				page = pdf[index]
				try:
					textpage = page.get_textpage()
					try:
						native.append({"page": index + 1, "text": textpage.get_text_range()})
					finally:
						textpage.close()
				finally:
					page.close()
			(output / "pdf-text-layer.json").write_text(json.dumps({"sourceHash": request["sourceHash"], "pages": native}, ensure_ascii=False), "utf-8")
		finally:
			pdf.close()
	class Capture(logging.Handler):
		def emit(self, record):
			nonlocal enrichment_failed
			if record.levelno >= logging.ERROR:
				enrichment_failed = True
			if record.levelno >= logging.WARNING and len(warnings) < 100:
				warnings.append(record.getMessage()[:1000])
	capture = Capture()
	logging.getLogger().addHandler(capture)
	try:
		result = (engine or converter(root)).convert(source, raises_on_error=False)
	finally:
		logging.getLogger().removeHandler(capture)
	state = str(getattr(result.status, "value", result.status))
	if state not in ("success", "partial_success"):
		raise RuntimeError("Docling conversion failed: " + state + "; " + "; ".join(str(error) for error in result.errors))
	for error in result.errors:
		warnings.append(str(error))
	material_defect = bool(result.errors)
	document = result.document
	if source.suffix.lower() == ".pdf":
		from native_code import recover_native_code
		recovery = recover_native_code(document, source)
		native_verified_ids = {item["id"] for item in recovery["recoveries"]}
		warnings.extend(recovery["warnings"])
		material_defect = material_defect or recovery["needs_review"]
		(output / "native-code.json").write_text(json.dumps({"sourceHash": request["sourceHash"], **recovery}, ensure_ascii=False, indent=2), "utf-8")
	if expected_pages is not None and len(result.pages) != expected_pages:
		warnings.append("Incomplete PDF page coverage: expected " + str(expected_pages) + ", received " + str(len(result.pages)))
		material_defect = True
	document.save_as_markdown(output / "document.md", artifacts_dir=assets, image_mode=ImageRefMode.REFERENCED)
	document.save_as_json(output / "structure.json", artifacts_dir=assets, image_mode=ImageRefMode.REFERENCED)
	# Staging directories are renamed by the host. Never persist absolute staging
	# URIs even if a future serializer changes its reference-path default.
	markdown_path = output / "document.md"
	markdown = markdown_path.read_text("utf-8").replace(output.as_uri() + "/", "").replace(str(output) + os.sep, "")
	markdown_path.write_text(markdown, "utf-8")
	def relative_images(value):
		if isinstance(value, dict):
			for key, child in value.items():
				if key in ("uri", "url") and isinstance(child, str):
					value[key] = child.replace(output.as_uri() + "/", "").replace(str(output) + os.sep, "")
				else:
					relative_images(child)
		elif isinstance(value, list):
			for child in value:
				relative_images(child)
	structure = json.loads((output / "structure.json").read_text("utf-8"))
	relative_images(structure)
	(output / "structure.json").write_text(json.dumps(structure, ensure_ascii=False, indent=2), "utf-8")
	blocks = extract_blocks(document, source.suffix.lower())
	for block in blocks:
		if "code" in block.get("kind", "") and len(block["text"]) >= 6000 and block["id"] not in native_verified_ids:
			warnings.append("Long recognized code block may reach the parser token limit; compare " + block["id"] + " with the original PDF.")
			material_defect = True
		if source.suffix.lower() == ".pdf" and "code" not in block.get("kind", "") and re.match(r"^\s*(?:(?:async\s+)?def\s+\w+\([^)]*\)\s*:|(?:export\s+)?(?:async\s+)?function\s+\w*\s*\()", block["text"]):
			warnings.append("Possible code region was parsed as plain text: " + block["id"] + "; compare the original PDF for line breaks and indentation.")
			material_defect = True
	formula_metadata = []
	if source.suffix.lower() == ".xlsx":
		blocks, formula_metadata = excel_provenance(source, warnings)
	for index, table in enumerate(document.tables):
		# Preserve merged cells in the native JSON and an HTML companion; GFM cannot
		# represent row/column spans. CSV is provided for machine validation.
		table.export_to_dataframe(doc=document).to_csv(assets / ("table-" + str(index + 1) + ".csv"), index=False)
		(assets / ("table-" + str(index + 1) + ".html")).write_text(table.export_to_html(doc=document), encoding="utf-8")
		cells = getattr(table.data, "table_cells", [])
		if any(getattr(cell, "row_span", 1) > 1 or getattr(cell, "col_span", 1) > 1 for cell in cells):
			warnings.append("Table " + str(index + 1) + " contains merged cells; Markdown is flattened. Refer to structure.json and table HTML/CSV.")
			material_defect = True
		try:
			picture = table.get_image(document)
			if picture:
				picture.save(assets / ("table-" + str(index + 1) + ".png"))
		except Exception:
			warnings.append("Table " + str(index + 1) + " has no source image crop.")
	if source.suffix.lower() == ".pdf":
		warnings.append("OCR and code recognition require review; code enrichment covers detected regions only and long code blocks can be truncated.")
	if not blocks:
		warnings.append("No readable text blocks were extracted.")
	status = "partial" if state == "partial_success" or not blocks or enrichment_failed or material_defect else "complete"
	if any(warning.startswith("Missing formula cache:") for warning in warnings):
		status = "partial"
	(output / "source-map.json").write_text(json.dumps({"schemaVersion": 1, "source": str(source), "sourceHash": request["sourceHash"], "blocks": blocks, "formulas": formula_metadata}, ensure_ascii=False, indent=2, default=str), "utf-8")
	return {"status": status, "warnings": list(dict.fromkeys(warnings))}


def smoke(root):
	from fixtures import make_fixtures
	folder = Path(tempfile.mkdtemp(prefix="qualification-", dir=root))
	try:
		files = make_fixtures(folder)
		engine = converter(root)
		for source, expected in files:
			print("Offline fixture: " + source.name, file=sys.stderr, flush=True)
			output = folder / (source.stem + "-output")
			response = convert(root, {"inputPath": str(source), "outputDir": str(output), "sourceHash": sha256(source)}, engine)
			text = (output / "document.md").read_text("utf-8")
			structure = json.loads((output / "structure.json").read_text("utf-8"))
			source_map = json.loads((output / "source-map.json").read_text("utf-8"))
			known_short_code_limit = source.name == "digital-short.pdf" and response["status"] == "partial" and any(warning.startswith("Possible code region was parsed as plain text:") for warning in response["warnings"])
			ambiguous_code_is_partial = source.name == "digital-ambiguous.pdf" and response["status"] == "partial" and any(warning.startswith("Native code layout could not be verified") for warning in response["warnings"])
			if source.name == "digital-ambiguous.pdf" and not ambiguous_code_is_partial:
				raise RuntimeError("Ambiguous PDF code characters were not reported for review")
			if (response["status"] != "complete" and not known_short_code_limit and not ambiguous_code_is_partial) or expected not in text:
				raise RuntimeError("Offline conversion fixture failed: " + source.name)
			if source.name in ("digital.pdf", "digital-short.pdf") and not known_short_code_limit:
				for marker in ("公司技术资料", "Name", "Value", "calculate_total", "return"):
					if marker not in text:
						raise RuntimeError("PDF structure fixture omitted " + marker)
				if not structure.get("tables"):
					raise RuntimeError("PDF table structure was not detected")
				if "```" not in text:
					raise RuntimeError("PDF code block was not recognized; this profile has not passed code fidelity qualification")
				if "def calculate_total(values):\n    return sum(values)" not in text:
					raise RuntimeError("PDF code text or indentation did not survive recognition")
			if source.name == "scanned.pdf":
				for marker in ("Name", "Value", "Total", "123"):
					if marker not in text:
						raise RuntimeError("OCR table fixture omitted " + marker)
				if not structure.get("tables"):
					raise RuntimeError("OCR table structure was not detected")
				if not all(marker in re.sub(r"\s+", "", text) for marker in ("公司资料", "扫描表格")):
					raise RuntimeError("Chinese OCR text was not recognized")
			if known_short_code_limit:
				native = json.loads((output / "pdf-text-layer.json").read_text("utf-8"))
				if not all(marker in native["pages"][0]["text"] for marker in ("def calculate_total(values):", "return sum(values)")):
					raise RuntimeError("Unrecognized short code lost its original text-layer evidence")
			if source.suffix == ".pdf":
				expected_cells = {(0, 0): "Name", (0, 1): "Value", (1, 0): "Total", (1, 1): expected}
				grids = [{(cell["start_row_offset_idx"], cell["start_col_offset_idx"]): cell["text"].strip() for cell in table["data"]["table_cells"]} for table in structure.get("tables", [])]
				if not any(all(grid.get(position) == value for position, value in expected_cells.items()) for grid in grids):
					raise RuntimeError("PDF fixture table cell values or row/column positions are incorrect")
			if source.suffix == ".pdf" and not all(block["locator"].get("page") == 1 for block in source_map["blocks"]):
				raise RuntimeError("PDF source blocks lost page locators")
			if source.suffix == ".xlsx" and not all(block["locator"].get("sheet") and block["locator"].get("cell") for block in source_map["blocks"]):
				raise RuntimeError("Spreadsheet source blocks lost cell locators")
			if source.suffix == ".pptx" and not all(block["locator"].get("slide") == 1 for block in source_map["blocks"]):
				raise RuntimeError("Presentation source blocks lost slide locators")
	except Exception as error:
		raise RuntimeError(str(error) + "; qualification artifacts retained at " + str(folder)) from error
	shutil.rmtree(folder)
	return [source.name for source, expected in files]


def setup(root):
	if importlib.metadata.version("docling") != VERSION:
		raise RuntimeError("Unexpected Docling version")
	(root / "ready.json").unlink(missing_ok=True)
	from docling.utils.model_downloader import download_models
	from huggingface_hub import snapshot_download
	def fetch_model(repo):
		print("Downloading CPU profile model: " + repo, file=sys.stderr, flush=True)
		# The generic downloader includes both ONNX/Torch layout and fast/accurate
		# tables. This profile uses the Torch layout and accurate table model only.
		patterns = ["model_artifacts/tableformer/accurate/*", "README.md", "config.json"] if repo.endswith("/docling-models") else None
		snapshot_download(repo_id=repo, revision=MODEL_REVISIONS[repo], local_dir=root / "models" / repo.replace("/", "--"), allow_patterns=patterns)
		print("Downloaded CPU profile model: " + repo, file=sys.stderr, flush=True)
	def fetch_ocr():
		print("Downloading Chinese/English OCR models", file=sys.stderr, flush=True)
		download_models(output_dir=root / "models", with_layout=False, with_tableformer=False, with_code_formula=False, with_picture_classifier=False, rapidocr_models=["onnxruntime:ch"], rapidocr_model_size="small", progress=True)
	def report_failure(future):
		if future.exception():
			print("Model setup failed: " + str(future.exception()), file=sys.stderr, flush=True)
	with ThreadPoolExecutor(max_workers=4) as pool:
		futures = [pool.submit(fetch_model, repo) for repo in MODEL_REVISIONS] + [pool.submit(fetch_ocr)]
		for future in futures:
			future.add_done_callback(report_failure)
		for future in futures:
			future.result()
	models = {}
	for path in sorted((root / "models").rglob("*")):
		if path.is_file() and ".cache" not in path.parts:
			models[str(path.relative_to(root / "models"))] = {"sha256": sha256(path), "size": path.stat().st_size}
	offline()
	passed = smoke(root)
	if parser_assets() != LOADED_ASSET_HASHES:
		raise RuntimeError("Parser implementation changed during setup; retry qualification")
	packages = {dist.metadata["Name"]: dist.version for dist in importlib.metadata.distributions()}
	receipt = {"profile": PROFILE, "docling": VERSION, "python": platform.python_version(), "platform": platform.platform(), "parserAssets": LOADED_ASSET_HASHES, "modelRevisions": MODEL_REVISIONS, "models": models, "packages": packages, "smokePassed": passed}
	(root / "ready.json").write_text(json.dumps(receipt, indent=2, sort_keys=True), "utf-8")
	return doctor(root)


def main():
	action, root = sys.argv[1], Path(sys.argv[2])
	if action == "setup":
		with redirect_stdout(sys.stderr):
			response = setup(root)
		emit(response)
	else:
		offline()
		if action == "doctor":
			emit(doctor(root))
		elif action == "convert":
			request = json.load(sys.stdin)
			with redirect_stdout(sys.stderr):
				response = convert(root, request)
			emit(response)
		elif action == "serve":
			engine = None
			for line in sys.stdin:
				try:
					with redirect_stdout(sys.stderr):
						if engine is None:
							engine = converter(root)
						response = convert(root, json.loads(line), engine)
					emit(response)
				except Exception as error:
					emit({"error": str(error)})
		elif action == "smoke":
			with redirect_stdout(sys.stderr):
				passed = smoke(root)
			emit({"passed": passed})
		else:
			raise ValueError("Unknown operation")


if __name__ == "__main__":
	main()
