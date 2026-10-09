"""Small synthetic fixtures; these verify engine wiring, not real-document fidelity."""
from pathlib import Path
from base64 import b64decode
from io import BytesIO

# Self-authored raster caption keeps the Chinese OCR check independent of OS fonts.
CHINESE_CAPTION_PNG = (
	"iVBORw0KGgoAAAANSUhEUgAAAa4AAAA8CAIAAACxTskxAAAegUlEQVR42u1dZ1gUV9uepYgsiIii2AVsoIIVFSwgIMUCEV+NEWKi"
	"gmBDMKBIFLGESwFRQBGNSoIGX5RLUQELFlBiFBSVKgLSO7vALkvb8v04n+eazO7Ozi6L0Tfn/sE1zJw5c+bszD3PeSpNIBBgCAgI"
	"CP9uKKApQEBAQFBCU4CA8OVj69atysrKbm5uU6ZMQbOBpMJ/Ej/99JOjo2NsbGx+fj6Px5Njz/v27Zs5c+by5cudnJx+/fVXqc69"
	"cuXKuHHjFixYYGtru3r16qqqKnmNau/evQkJCQ0NDT09PZ9hegUCgb6+vpmZ2dKlS+3t7W/evCmyWVJSEo1Go9FoampqWVlZ+EOL"
	"Fi0Ch1asWCHfsf3888/6+voLFixYtmyZr6+vvLotKChQVVXV0tJSU1MbNWpUZ2enuJY8Hi8+Pj48PHzq1KkjR468evUqeh/75BFE"
	"kIi2trYBAwbASfvll1/k2HlERAToduTIkc3NzVKdm5OTA0cVHBwsx1F5eHjAnv/zn/98hkmeNGkSuNycOXNaWlpEtnn9+jVoY2pq"
	"Sji0bt06cCg5OVm+AwsNDYVT8fr1a3l1i+c+8t/u3r17sKWFhUV7ezt6JeUO6aRCNptNk4Tly5fLzMuOjo40eYPJZJJc8cGDB5aW"
	"lmlpaeQDO336NIvFAtuRkZF+fn5y/Bppa2uDDTMzMy0tLXHNcnNz16xZ09XVhd85atQouL148WI5jmrIkCFgQ09P77fffvsMX+XB"
	"gweDjR07dgwcOFBkG01NTbBBo9EIhzQ0NMQd6iUGDRoENuh0+vTp0+XVrYqKirq6OtjW0dEhafnHH3/AkcTFxdHpdCTDfQUL5D//"
	"/PPLsUorKirC51gYNTU1zs7Ojx49Mjc3Nzc3F0eIZWVlhw8fBtt79uzZtm2byGZ8Pr+np6ejo4PFYjGZzIaGBnEs3NLSEhYWBoUC"
	"+GQrKyuDDQ6HM3PmzLNnz5aUlIA9Fy9eNDExuXbt2r59+/BdDRw4EL758qUASEZTpkxRVVUlHK2oqHj8+HFhYWFNTU1bW1tXVxeP"
	"x4O/u0Ag6O7uZrFYdXV1BQUFqampFy5cIIyc5Ip4AZwANTU1cYeEBykvQMJSVVWV7ySTPJwQTCYzPj4ebB88eHDYsGGItv6BBXJ7"
	"e7uFhUVcXFxXV5dAIICSETnevXsnm4zq4OAg37uj0+kkl3N1dSW0t7e3LywsxLfh8XjW1tayXf3MmTPCF/X29gbv89WrV8GeBw8e"
	"gPbr168He/BasMuXLwsEgpMnT/7/t0tB4fnz5/gOITtkZmbKcb1w+vRp0K2Dg4Pw0djYWGlnQ09Pj/yKq1atAi1v3LgBdzY0NKSn"
	"pxcVFTU1NXE4nObmZihBE07fs2cPOJSSksLj8fBcnJGRQbIe7yP4+/tT1wnExsaKaxMSEkL9ovr6+nw+H612ZYAEKoTaa3V1dWtr"
	"awaDIa4ln88fMWIEaHzgwIF/9q78/f2hEEHSjMvlRkdHE9YmdDq9trYWtpF5LdyvX7+mpibhi65evRrSLtjz9OlTAhWGh4eDPebm"
	"5nB6LSwsoJjW09MDO4TClHyp8MKFC/KlwtmzZ5Nfcf369cJUmJKS0ntWGjt27JdJhdOmTSOnQj6fP378eOoXDQ0NRaTWJ7pCSIVd"
	"XV3btm0jkedpNJq9vT3Y/losXIqKim5ubh8+fPD29oYLn4CAAEiOERERQUFBJIs1AwMDXV3dESNGaGlp0el0RUVFeNTOzg4qv/DY"
	"vHkz2Lh//35TU5PInh8+fAjNuHB6o6Ki+vXrh2FYXl4eFNkwDFNS6hOPKLl3K1HDBfUD/x5InOTExMTi4mKKvQ0YMGDTpk1opSt/"
	"XSGLxbp27RrY9vPzk7h6/eabb8BGUVHR3bt3v5YpUFdXDw0NTU5OHjp0qJ2dnY+PD9h/8uTJnTt3Yp/sldBsV1hYCHULT58+LS0t"
	"ra6ubm5ubm9vf/78OezWxcVF5OWsra2BrYPL5YoUedhs9v379zEMGz16NH5tPmnSJDiewMBABoMhx0l4/fr127dvq6qqWltbu7u7"
	"+Xw+XikGFpuNjY0lJSUZGRlsNtvCwuLp06cFBQWVlZWNjY0tLS0cDqerq6unp4fL5fJxsLGxgWKyxC/T5/zdo6Ojnzx5kpWVVVBQ"
	"AH/EtrY2cCNcLpfH44FbuHHjBjTsgD08Hq+np6e7u5vD4bBYLAaDUV9fX1FRUVRUlJOT8/z58wcPHgg/AF1dXUlJSW/evKmrq2Oz"
	"2VwuF04yn89vbW2trKzMzs6+c+cOl8sFTwh+UfLy5UuR4szatWtBA1dXV3HmJoRe6QrPnj0L2owYMYLNZkuUMLlc7siRI7FP1swv"
	"f4FMQG1tbWNjI9gODg6GUzR9+nQmkwmb7dixAx7av38/vgdbW1vsk6Gzs7NT3IUgo61Zs0Z4gRwTEwP+9fHxIZzIYDCgYP7TTz+B"
	"nVD27M0CGRIWRs0hToaebWxsCIfS09Nfv35dWVnZ0tLS3d3t5uYGWl6/fp3BYBQXFz98+LC6ujojI6OoqKiurg7Yo8TpCuGPnpyc"
	"3NPTAxSLpaWlmZmZ9+/f783jhKfC3vTT2NhIcYZZLJZAIIiKisLv/P7774X7LCgoUFBQACIhfHoRZIBY+ZzL5UI68PPzI7Hc4b/q"
	"mzdvDgwMxDAsLS3t5s2bjo6OVH749vb2pKQkTU1NTU1NdXV1dXV1Op3er18/ZWVlZWVlBQUF8GPjbbXAXAue+La2tubm5vr6emtr"
	"ayrjFAe80tDLy2vatGk7d+5UVVVNTU2FPhwfPnyIjo4Gyz0OhxMcHLxu3ToDAwNwF5MnT1ZWVi4pKbGyslJRURF3IScnp/DwcHV1"
	"deiwggekwqVLlwobHHfu3BkYGEin0+FX5+uFl5fXq1evhPdDdSogBVNTU7zILLFbGo2mpKSkpKQEHJh1dXVJGpeUlOTn52tqampo"
	"aIBnr3///uDBU1RUBP5YIgUIIBt2d3eDh5DNZrNYrJaWltbWVicnp95PDpvNPnjwIH7P1atXDx06NHbsWPxOT09PPp+PYdju3btF"
	"Pk4IvZUKL168CEXCjo4Oisza3NwMPSFGjRrV1tZG5ayPHz/K5V4+fvzYG6kQuIAEBQUdPHgQ/Mtms/E+z11dXbNnzwY9JyQkAH9A"
	"IyMjvMxIBTwe7/r16xwOR5zZpKys7Pz587ABQbJwcnKCd/pVS4WzZs2iKB9BQB8GEqkwJSWF+vAIklefvlNSSYXQxZ1Op0OP0R9+"
	"+AHfIfT3HD16NGGiEORjQWYymdB96dSpU1L1+PPPP8NfdOXKlVRM+18IFT569AhY9IRfM7D8h197W1tbfBCCiYlJfX29zL+BMBVK"
	"BblQYWZmZnZ2dkVFBZPJ7Ozs5HK50Ebs4ODQ2dnJZDJramoKCwvT0tKkeusQFeLDS27evPnXX38VFxc3Nzd3dnYaGxuDU3777Tc2"
	"m11TU/Pu3bt79+7h4w5jY2PfvXsHhFMajfbs2TPQW35+PnB4pNFoqampiMv6hAqh1kZXV1dYJKypqSF3RcTL8N7e3l8+Fb548cLK"
	"ygq/wiovLyd8G6AeUEdHp6qqCpCjiYkJ3JmYmEhFar5z505mZmZFRUVLSwvwTMZTITBQ1NfXl5aWvnr1KiUlBSqAeDyeSIdNuVAh"
	"ibuMSGcauVDhvXv3nj179uHDh8bGRg6Hs2XLFtDy2rVrDQ0N+fn5SUlJeLchILaHhYWFhYXFx8cLax6DgoKCgoJKS0u/TCokoLq6"
	"Gmp+8M40ubm5MHLG3d0d7ISm4YkTJ7a2tjY1NUFnIE9PT0RkfUKF0GoMFNiEoz09PYMHD161alVWVpa4TpOSkvBPhqurK5fLJRkE"
	"m82+fv36/fv3X7x4kZeXV1ZWBo2SeIskUNDw+Xwul9vd3d3e3s5gMCorKwsLC58/f56YmAhtOxSpkM/n37p1y9zcHD/aiRMnEoJY"
	"k5OTIblraGi8ePEC/zSPGTMGnjt37tyYmBhx8bNA+JL2vXr69Ck4cfbs2erq6gSO/qqpkABIhXi/QiogBPlUVVW9f/8+MzPzwYMH"
	"cgkZlpfZhIDjx4/jRT+4v6OjIyEhwcLCYtq0aVAQYTAY8CG0s7ODipolS5aA8AcEOVNhTk4O/CKJtALjaW7WrFkiAyoEAsHu3bvx"
	"77OFhYXwO/wPWpDLysoCAwMJCnU1NbVjx47hH6zMzEw7Ozu8XUXYoaG0tJSQN0lZWXnWrFnu7u4nTpwgRIbIRoWdnZ3Dhw8XR0z/"
	"81T4/v17Jycnc3PzGTNm6OvrDx8+XFNTExg3CPY0TFTWny+TCvl8/oQJE/BUWFlZuX37dhC3ChdY+FNevnxJsAoaGxu3trYiFpM/"
	"FVZUVEC7ZL9+/XJycoRPgM6DYCEJLQzCH2pLS0vs79HyERERn+cLJpEKYTIYiFWrVlVUVMAGaWlpCxYsIPgDitMMtLW1ubq6Cpsa"
	"VVVVCXMos1SIT4jw8OHDf4oKwXL+M1NhQUGBzGvVY8eOiVRT3Lhx4/Hjx69evSoqKqqqqmpqampra+vo6ABLECpUCL0LOzs72Ww2"
	"EEjLy8sLCwuzs7OfPn2anJxMorUErqN4oznwuyT3QsMr4jEMW7duXUNDA2IxOVNhcXExXko6cuSIcOvKykp8SEBISIiANLEVFOMh"
	"xowZk5aW9o9TIZ/Ph0w3bty4pKQkQgMOhwPTLkyePPm///2vxItmZGTA2DiAw4cPE9q0tLSkpqYCAwWDweBwOOXl5TDkAOgK2Ww2"
	"0BVmZ2c/fPgQmLD5fL6RkRH2KYfVP0KFd+/eNTIywjOamZkZdVaSlgpbW1sPHjzY2dnZ0NAgMxWGhYXJRU0hGyZNmiTuZvG6aQzD"
	"IiIioMfozZs3hduXlpbipRC8hOHj4yOVTR+BjApfvnwJl2AYhllaWoJEIwTASAwMw+bNmyfROtzU1DRnzhz8L7d27dru7m5x7YFP"
	"ogyIiYmR1sX62bNnysrKvr6+JNnfoqOj7927J3IqxCErK8vNzW3IkCF0Op0kZBsiICAA3gW5BRnKJiAe63NSYXJy8vz588G/y5cv"
	"72sqjIqK8vf3HzRokIaGBvimimMWGLZM4IKwsDDsU0a1L5AKX7x4gf09GU9sbOzRo0ehuICPaMjPz3d1dcW7qYp0dTQ2Nt61a9eN"
	"GzdKS0tRRgYZqfDWrVv4iR43bpxI1xAmkwl/ORqN9tdff1G5AIvFgh/AqVOnki+Qpc3hjOHSDsoQbVJWVga3YeqX3gDSBI/HKy4u"
	"pqLvhzksIBXy+fzjx49XVlYKhLwRYWT+/Pnz+5oKf//9d5HxcFZWVn1NhXin/c7OTpiiUQYqjIqK+gKpEFjqjI2NoTtRbGxsfX09"
	"8Ot2d3dvaWnp6OiIj4+3tbUlEN+UKVMyMjJu3rwpMsIduiKKTMaDICBPx2BraxsYGAjSvWlqaiYlJQ0dOlR4fk+cONHa2gq2nZ2d"
	"586dSzHCNyUlxdPTE8OwM2fOkAeikiSqIwc09WBSJizBhPKn9gYw8aqCgoK+vr7E9teuXaupqcH+nspw5cqVvr6+0J8JQkFBAUwj"
	"hmGtra1QVpIvPn78GBERYWNjAzOYgeoF/fr1c3R0jIuLS0xMxGvi0tPTc3JySktLa2pqmpubWSwWIYaX3Hmby+W+evUqNDTU3t6e"
	"kO1m5syZe/fu5XK5EoOXMXnnyBF88qiHAT9KSkqPHz8GmhPqbxcMV8cjLi7uyZMn4CXC7x86dOjRo0cfPnxobm7u6empo6OzZs2a"
	"u3fvwiyQAwcOPHHixJs3b0xNTR0cHHJzc2H0sXDc1Lx581D8iIw+UC9evBg+fDjB4onXNEO6GTBgALlroUCMI7HENrdv36buD4hP"
	"nkiwTsjgYn39+vXeT6abm5tAmsgWQJdQtlq/fj14SeA7IxDyOhozZkxwcDBeyYCXCpubmxMTEz08PCi+7fjoF3d3d5H5oIYPH37q"
	"1CmRCcd6aTbZvHmzyHQ1Li4uhKyR8pUKKfIgPmw0Kiqqrq7O09Nz+PDh4tgwJiZGoh68ubkZyBnKysq1tbV4qVDw9/h0PPr37+/t"
	"7S3yJ3jy5MnChQsJ7eVbcOLfaEEmybng5eWFSarDwGAwEhISejMamKlFWiokvAkyUCGJw0RLS0tycjJJrDvM2bNlyxbqNwuTEp4/"
	"fx6/QP7xxx9hyKNw5CJBw9De3g7V7aNHj8ZntCbRyQqDJE993znTiCuZJOxX+JmpsLW1dcmSJXA8R48e7enpgc4V58+fF4jK7AkC"
	"n48fP06iqoOS4Lp16wS4eBtIhd3d3fj1lpqampeXV3V1NfmA09PTv/322+nTpy9btkxFRYWKcgaBLF+huFwGZWVlZ86cgeZUuEyD"
	"6OzsDA4O1tfXX716NZTs/meQl5dnb2+vra09aNCgWbNmlZaW9r5PBoMBigTMnz9/8uTJ+EMhISFgoV1TU3Po0CHCiRwOJy0tLSws"
	"zMXFZcqUKRoaGjBZC1AvQmG/vr6e+nhgFkWgKwgICMDn5ukjLFq0CIarBwcHf/fdd1/Cz52fnz9v3rxHjx6Bf4OCgvbt26ekpLRr"
	"1y7sU/Y2kAEB4tdff129ejWXy+Vyub6+vkuWLBFZy+H8+fOXL18G2yKlPyAtxsXFaWhozJgxw8DAoL29PSwsbOTIkeQFfBYtWnT1"
	"6tU5c+bcuXOHyWRSUc4gYDLUQfb29oaq68jISOEUmx8/fty/fz9o88MPP2RnZ+NjML52VFdXQ0Weg4ODnp5e7/vcsmULCM4X/q5o"
	"aWkFBAR4enrq6urOmjUrKysLZMHLzc199+5dRUUFRloexMDAQE9Pb+zYsSSpcTBRWRTBpf38/LZu3Uqn0+FL23cYN27c2LFj/fz8"
	"Nm7cqKys7O7uTt7+/fv34kqLgORAvVcWnTt3ztvbm8PhAKNNZGQkHNWOHTsiIyPLy8vz8vLi4uKAQNrW1rZz507o9TlmzJjt27fP"
	"nz+/f//+hM7fvn0L07uZm5uT6PJ0dXXfv3+vo6Pj6uoqlU8lWB/0XY2XfzsVpqamwsWjs7MzwXEaPoiBgYEg5TKDwVi3bl16evpn"
	"zsTZd/jw4QP2yWXa29u79x3GxMQA1eT48eNXr16NT/gKiZLFYnl5eW3atIk8JXj//v1BplUMw+7cubNs2TLZhjRx4kQfHx8/Pz8q"
	"hYfkaGD98OHDF5K8+s2bN9u3b8/IyIDS2a1bt2DgOYZhKioqhw8f/v777zEM27t3r6OjY2pq6rZt28CXUkdH58CBA5s2bRJn5Bk1"
	"atSgQYPq6uowDNu/fz9GIV+ctKT2L0wDjn22inccDmfr1q3wm4MvCIsJVUyHuQn+/PNPkjz4GLUMhhKrespsbpYW0Ats6dKlsPiZ"
	"zMjJyYGLI19fX5EfDBUVFX9/fzqdTkhOBzBhwgRnZ+fIyMjMzMy2tjaYS7GXtdCOHz/+OXkQmGWlentl0BVSQXZ29po1a2bOnAl5"
	"EMMwDQ0NPA9CUQD4wVRVVU2bNs3R0bG6ulpNTS0gIKC4uNjDw4PE2D148GCQC9nCwgKviCR7OT+FFZI7nAo7ISHIXyrcs2cPlImC"
	"g4NFOtlgn1zAoqKi5syZAySUwMBAW1tb4WgTcvT09Mi8tOnlRBC0Pxgul/2zZ88woZSiIgH8TkhQXl5ua2sLDD6GhobQQoKRuvvQ"
	"aDRjY2Nzc/NFixYtXLiQkKETrx9EDzR1NDY2JiYmxsTE4BlQXV2dJDssjUa7dOnStGnT2Gz2x48faTSai4tLUFAQ3jmUBA4ODosX"
	"L8ZnYUD4OqTCx48fw3JCNjY2EivIzJw5E36guFyui4sLLPVLER0dHbLdRnd3dy8nglBqHcNZM0EVEXV19ZUrVwKNQU5OjlSdQFHX"
	"xsYGOhKeOHFCYpWfhQsXXr58ubGxMTs7Oyws7JtvvhHOVAzqYJCwOQImVJR1xowZQBmH58FVq1a9fv0ak6TfhGsjZWXlFStWUORB"
	"7JMnqbTyAcI/TIUsFuvHH38EgoaGhgZ0+CAQUFNTE4iWTUtLu337tqGhIVRsFxYWHjhwQKrRQBduaZ1pxBVfl+HSBP0AVOt4eHho"
	"aGhUVFQ4OTkZGRlZWFjcuHGDIAaK7ATDGeihXOnk5EQlcfTUqVPXr19PEleAF0UlyqTyQlFRUWpqKrDn1NbWgtAIkIftq3jonZyc"
	"Wltb8V+OSZMmJSUlJSQk4MNPMfHeo+Cj2N3dvWbNmn379lFfzcjFkx/hcyyQi4uLGxoaWlpaYmJiysvLsU8Oa/7+/iDCAQBskwtB"
	"QPBxcnKiGJeCYZi4YpgSQYjZwGQqWwycvKDqjcfjbd68GegHtLW1Qfz15s2bQZjHkydPnjx5Mnbs2G3bti1YsGDcuHFAi0d+lSNH"
	"jjQ0NCQmJsoxYyhkQCge9jXi4+MlKv7lqMToDcWLvOiwYcOio6NBJMmIESP8/f1dXV0lai3LysrOnTtnaWlpaWkZFxdnaWkJYk+D"
	"goJu37598uRJkRZFihohpNz44qgwMjLy1KlTmJBjXV5enmwP8aZNm7Kzsylqx6XKQaKqqgpdsg0NDakvVEVi48aNGzduxHB+G25u"
	"bunp6UBDdOHCBfA9/+WXX4YNGxYfHw+W5OXl5b6+vnQ6fcOGDV5eXhKpEMOw8PDwDRs24KUD+KrLtsLtIyqE3YrzX5H5tf9sJ5JM"
	"qbW19XfffWdmZrZx40ZhxxfCPCQnJ589e/bu3buwQDudTk9KSlq0aBF4L3Jzc62srBYvXrxr1y57e3vqkYLw1sjp/sqVK1euXEFs"
	"9VmpcOfOneHh4VS+UUpKSkOHDtXW1oZ/tbS0tLS0Bg8eXFtbC51O8vLyQkJC8OVcMQoefBi1gu7CNj7w9b516xbF2rvCyMzMjIiI"
	"iIuLA8yioKAQHR29YsUK7FP4amxsbEhIyKlTp06fPg0kRA6HExUVBZr5+PiQJyno378/oQGkMBneeYFAAE8H3nC9oSpFRUVotXz5"
	"8iX2yZwthcJFQUFBQQGY+EEMMnyWyO+Oz+dDAx2BfOF9yeBXSHJRieRSW1t74cKFc+fOVVZWYrhkkeCTqaWllZGR4ezsfOfOHXAo"
	"LS0tLS1t4MCBVlZWpqamBgYGJiYm5MoNODyRlC2tqNh7jTmGYpDxwKfVotPphoaG9vb27u7uR44cuXTp0r1793JychoaGsgTAeEr"
	"+GhoaJDkuBfIIyGwmZnZ0KFDR48eTdDFGBgYUOmhrKwsISHB09OTEISrra1NklyeyWQeOnQIpmDAcLYOQkkAipHX+PxXFIGnvz/+"
	"+KM301hdXa2lpaWtra2np4d/gbdt20ZomZeXl5KSkpWVVVJSUl9fD0oviKvZAFWiJiYmhEMdHR1Tp07V09MzNDTEX5Ewe1BRIwNE"
	"pt2kooBWUFAgGLU0NTU9PDwIdSz4fP7BgwdFmr+WLl0qsXo4NLmcPn1a+CjMmEkR27dvR/Fz8qxtEh8fHxoampqaCooZyQaY3cDa"
	"2jovL6+vb0NcMoUdO3YISEv8GBkZifRPVFFRcXV1pZKDgMlk7t27l+AN++2331If/KVLlzAxVdwEFOpGwYuePXu2l9NIKLwL8Pvv"
	"v8slBllfX19kCgNMVOgbvs3bt29lpsLdu3dTHGdJSUl4eLiwvxSNRrOysoqLiyMpgZuXlwfCdTBc9kCJPCgQCOCzJzLbNtTYUPQr"
	"dHFxQYwm54p3vQeXyzU3NxeZj7cvwOFwhL/MkyZNqqurE5CmCxR2dp0wYUJAQAD5iSJFqg0bNoDl2/jx4ykWgAY4duyYVDKsQEyx"
	"wKNHj/ZyGoWj+tauXUteoos6FQ4cOFBkeh68IK+kpLR3715Cm9TUVIlJoQmALtbAC4IK3r59S3iEBg8e7OPjU1JSQrGHpKQkkMZc"
	"UVHx7du3AgrZKuG1hO9aIBCsWrVKKirEp5JEkBZKfbTuVlRUfPz48Wdb5quqqhoYGOTk5CgpKeno6EydOnXZsmUbN24UmQMK/+Jd"
	"vHjRxMRkzJgxxsbGc+fONTc3p2L6EMaIESNiYmJcXV1379595swZqSJhhgwZAvhC2GcQo1AmuPf2d4jRo0cbGhoqKioaGBjMmDHD"
	"zs4OFuqVGRMnTgQDA9pDQnQNcM179+7dnDlzTE1NbWxshN1NQLCaVBgwYABIJCOxDhSEkZGRh4cHKHpjaGjo7e29fv16cosKAfb2"
	"9vb29rm5uenp6bD6AkaakgNu46V7iKqqKirXnTJlCnh+RMYmIVAE7X/Jii8QCORl7kT4F6K4uNjZ2Xnfvn0rVqxADxKiQgQEBAQM"
	"RZsgICAgICpEQEBAQFSIgICAgKgQAQEBAVEhAgICAqJCBAQEBESFCAgICIgKERAQEBAVIiAgICAqREBAQEBUiICAgICoEAEBAQFR"
	"IQICAgKiQgQEBAREhQgICAiIChEQEBAQFSIgICAgKkRTgICAgICoEAEBAQH7PzsHvqY55zH8AAAAAElFTkSuQmCC"
)


def make_fixtures(root):
	from PIL import Image, ImageDraw
	from docx import Document
	from openpyxl import Workbook
	from pptx import Presentation
	from pptx.util import Inches
	files = []
	# Raster PDF exercises the actual OCR path without a hidden text layer.
	image = Image.new("RGB", (1000, 600), "white")
	draw = ImageDraw.Draw(image)
	draw.text((80, 70), "TOTAL 123", fill="black", font_size=60)
	draw.text((80, 180), "Document conversion fixture", fill="black", font_size=32)
	image.paste(Image.open(BytesIO(b64decode(CHINESE_CAPTION_PNG))), (80, 260))
	draw.rectangle((75, 350, 700, 520), outline="black", width=3)
	draw.line((75, 430, 700, 430), fill="black", width=3)
	draw.line((400, 350, 400, 520), fill="black", width=3)
	draw.text((95, 375), "Name", fill="black", font_size=32)
	draw.text((425, 375), "Value", fill="black", font_size=32)
	draw.text((95, 455), "Total", fill="black", font_size=32)
	draw.text((425, 455), "123", fill="black", font_size=32)
	path = root / "scanned.pdf"
	image.save(path, "PDF", resolution=120)
	files.append((path, "123"))
	# A born-digital PDF preserves text independently of OCR.
	objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R /F2 6 0 R /F3 7 0 R >> >> /Contents 5 0 R >>", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
	chinese = "公司技术资料".encode("utf-16-be").hex().encode()
	stream = b"BT /F1 18 Tf 72 720 Td (Document conversion fixture 456) Tj ET\n"
	stream += b"BT /F3 16 Tf 72 680 Td <" + chinese + b"> Tj ET\n"
	stream += b"72 500 400 100 re 72 550 m 472 550 l 272 500 m 272 600 l S\n"
	stream += b"BT /F1 14 Tf 85 570 Td (Name) Tj 200 0 Td (Value) Tj ET\nBT /F1 14 Tf 85 520 Td (Total) Tj 200 0 Td (456) Tj ET\n"
	header_stream = stream
	code_lines = ["def calculate_total(values):", "    return sum(values)", "", "def print_report(values):", "    total = calculate_total(values)", "    print('Total:', total)", "", "print_report([1, 2, 3])"]
	stream += b"0.95 g 62 284 470 164 re f 0 g\nBT /F2 12 Tf 72 430 Td\n"
	for index, line in enumerate(code_lines):
		if index:
			stream += b"0 -18 Td\n"
		escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
		stream += b"(" + escaped.encode("ascii") + b") Tj\n"
	stream += b"ET\n"
	objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
	objects.extend([b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>", b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light /Encoding /UniGB-UCS2-H /DescendantFonts [8 0 R] >>", b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light /CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 4 >> /DW 1000 >>"])
	def write_pdf(name, content):
		objects[4] = b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"
		data, offsets = bytearray(b"%PDF-1.4\n"), [0]
		for index, value in enumerate(objects, 1):
			offsets.append(len(data))
			data.extend(str(index).encode() + b" 0 obj\n" + value + b"\nendobj\n")
		xref = len(data)
		data.extend(("xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)).encode())
		for offset in offsets[1:]:
			data.extend(("%010d 00000 n \n" % offset).encode())
		data.extend(("trailer\n<< /Root 1 0 R /Size %d >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)).encode())
		path = root / name
		path.write_bytes(data)
		files.append((path, "456"))
	objects[5] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier /Encoding /WinAnsiEncoding >>"
	write_pdf("digital.pdf", stream)
	# StandardEncoding maps this font's quote byte to U+2019. Retain the
	# ambiguous original case: a recognizer must not silently turn it into ASCII.
	objects[5] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>"
	write_pdf("digital-ambiguous.pdf", stream)
	# Keep the short, unboxed snippet which the layout model misclassified in the
	# first real qualification: this must be preserved or explicitly report partial.
	write_pdf("digital-short.pdf", header_stream + b"BT /F2 12 Tf 72 430 Td (def calculate_total\\(values\\):) Tj 0 -18 Td (    return sum\\(values\\)) Tj ET\n")
	document = Document()
	document.add_heading("公司资料 Fixture document", 1)
	document.add_paragraph("Office text 789")
	table = document.add_table(rows=2, cols=2)
	for cell, value in zip([cell for row in table.rows for cell in row.cells], ["Name", "Value", "Total", "789"]):
		cell.text = value
	path = root / "office.docx"
	document.save(path)
	files.append((path, "789"))
	book = Workbook()
	book.active.append(["Name", "Value"])
	book.active.append(["Total", 321])
	path = root / "sheet.xlsx"
	book.save(path)
	files.append((path, "321"))
	presentation = Presentation()
	slide = presentation.slides.add_slide(presentation.slide_layouts[6])
	slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(2)).text = "Presentation fixture 654"
	path = root / "slides.pptx"
	presentation.save(path)
	files.append((path, "654"))
	return files
